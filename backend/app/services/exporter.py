from io import BytesIO
from openpyxl import Workbook
def tender_xlsx(p):
    wb=Workbook(); ws=wb.active; ws.title='Tender Summary'; ws.append(['Field','Value','Source Page','Confidence'])
    labels={'tender_title':'Tender Title','reference_number':'Reference Number','closing_date':'Closing Date','closing_time':'Closing Time','pre_bid_date':'Pre-Bid Date','pre_bid_time':'Pre-Bid Time','pre_bid_location':'Pre-Bid Location','submission_location':'Submission Location','delivery_period':'Delivery Period','bid_validity':'Bid Validity','quantity':'Quantity','warranty':'Warranty','bid_security':'Bid Security','performance_security':'Performance Security','required_documents':'Required Documents'}
    for k,label in labels.items():
        f=p['summary'].get(k,{}); ws.append([label,f.get('value'),f.get('page'),f.get('confidence')])
    r=wb.create_sheet('Requirements'); r.append(['Category','Requirement','Page','Confidence']); [r.append([x.get('category'),x.get('requirement'),x.get('page'),x.get('confidence')]) for x in p.get('requirements',[])]
    z=wb.create_sheet('Risks'); z.append(['Risk']); [z.append([x]) for x in p.get('risks',[])]
    for s in wb.worksheets: s.freeze_panes='A2'
    out=BytesIO(); wb.save(out); out.seek(0); return out
